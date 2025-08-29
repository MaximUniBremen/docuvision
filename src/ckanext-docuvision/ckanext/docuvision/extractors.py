import logging
import os
import subprocess
import PyPDF2
from pdf2image import convert_from_path
from docx import Document
import openpyxl
from PIL import Image
import pytesseract

# Create a logger for logging information and errors
log = logging.getLogger(__name__)


def extract_text_pdf(filepath):
    text = ""
    try:
        # Open PDF in read-binary mode
        with open(filepath, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            # Loop through each page to extract text
            for page in pdf_reader.pages:
                text += (page.extract_text() or "") + "\n"
        if len(text ) <5:
            text = extract_text_tesseract(filepath)
        return text
    except Exception as e:
        log.error(f"Error extracting text from PDF: {str(e)}")
        raise

def extract_text_tesseract(filepath):
    """
    Reads the PDF at the given filepath and returns extracted text as a string.
    """
    try:
        # Convert PDF pages to images
        images = convert_from_path(filepath, poppler_path="/usr/bin")
        extracted_text = []

        for i, image in enumerate(images):
            # Convert PIL image to RGB just to be sure
            img_rgb = image.convert("RGB")

            # Use pytesseract to extract text
            text = pytesseract.image_to_string(img_rgb)
            extracted_text.append(text)

        return "\n".join(extracted_text)

    except Exception as e:
        print(f"Error extracting text: {e}")
        return ""

def extract_text_docx(self, filepath):
    try:
        doc = Document(filepath)
        text = "\n".join([p.text for p in doc.paragraphs])
        return text
    except Exception as e:
        log.error(f"Error extracting text from DOCX: {str(e)}")
        raise

def extract_text_doc(filepath):
    """
    Robustly extract text from .doc (Word 97-2003) files.
    First try textract (requires antiword/catdoc), then fallback to antiword directly.
    """
    try:
        import textract

        try:
            text_bytes = textract.process(filepath)
            text = text_bytes.decode("utf-8", errors="replace")
            return text
        except Exception as textract_err:
            log.warning(
                f"Textract failed for DOC: {textract_err}, trying antiword fallback."
            )
            # Try antiword directly if textract fails
            try:
                output = subprocess.check_output(["antiword", filepath])
                return output.decode("utf-8", errors="replace")
            except Exception as antiword_err:
                log.error(
                    f"Both textract and antiword failed for DOC: {antiword_err}"
                )
                raise Exception(
                    f"Failed to extract DOC text: textract error: {textract_err}, antiword error: {antiword_err}"
                )
    except Exception as e:
        log.error(f"Error extracting text from DOC: {str(e)}")
        raise

def extract_text_xlsx(filepath):
    # Try openpyxl first (for .xlsx and compatible)
    try:
        log.info(f"Attempting to read file: {filepath}")
        if not os.path.isfile(filepath):
            log.error(f"File does not exist: {filepath}")
            raise Exception(f"File does not exist: {filepath}")
        size = os.path.getsize(filepath)
        log.info(f"File size: {size} bytes")
        if size == 0:
            log.error("File size is zero bytes, cannot extract.")
            raise Exception(f"File is empty: {filepath}")
        wb = openpyxl.load_workbook(filepath, data_only=True)
        text = ""
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                text += (
                        "\t".join(
                            [str(cell) if cell is not None else "" for cell in row]
                        )
                        + "\n"
                )
        return text
    except Exception as e:
        log.warning(
            f"openpyxl failed on {filepath} with error: {e}, trying xlrd as fallback."
        )
    # Fallback to xlrd if openpyxl fails (maybe it's actually an .xls file or old Excel format)
    try:
        import xlrd

        wb = xlrd.open_workbook(filepath)
        text = ""
        for sheet in wb.sheets():
            for row_idx in range(sheet.nrows):
                row = sheet.row_values(row_idx)
                text += (
                        "\t".join(
                            [str(cell) if cell is not None else "" for cell in row]
                        )
                        + "\n"
                )
        return text
    except Exception as e2:
        log.error(
            f"Both openpyxl and xlrd failed to read Excel file {filepath}: {e2}"
        )
        raise Exception(
            f"Failed to extract text from Excel file. openpyxl error: {e}, xlrd error: {e2}"
        )

def extract_text_image(filepath):
    try:
        log.info(f"Attempting OCR on image file: {filepath}")
        if not os.path.isfile(filepath):
            log.error(f"Image file does not exist: {filepath}")
            raise Exception(f"Image file does not exist: {filepath}")
        size = os.path.getsize(filepath)
        log.info(f"Image file size: {size} bytes")
        if size == 0:
            log.error("Image file is empty, cannot extract.")
            raise Exception(f"Image file is empty: {filepath}")
        image = Image.open(filepath)
        text = pytesseract.image_to_string(image)
        return text
    except pytesseract.pytesseract.TesseractNotFoundError as e:
        log.error(f"Tesseract is not installed or not in PATH: {str(e)}")
        raise Exception("Tesseract OCR is not installed or not in system PATH.")
    except pytesseract.pytesseract.TesseractError as e:
        log.error(f"Tesseract OCR failed: {str(e)}")
        raise Exception(f"Tesseract OCR failed: {str(e)}")
    except Exception as e:
        log.error(f"Error extracting text from image: {str(e)}")
        raise
